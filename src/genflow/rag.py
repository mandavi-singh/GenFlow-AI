from pathlib import Path

from langchain_core.documents import Document
from langchain_core.vectorstores import InMemoryVectorStore
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from genflow.config import settings
from genflow.llm import ainvoke_with_retry, get_llm

DOCS_DIR = Path(__file__).resolve().parents[2] / "data" / "docs"
INDEX_PATH = Path(__file__).resolve().parents[2] / "data" / "index.json"
CHUNK_SIZE = 800
CHUNK_OVERLAP = 100

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP, separators=["\n\n", "\n", ". ", " "]
)
_embeddings = None
_store: InMemoryVectorStore | None = None


def _get_embeddings() -> OllamaEmbeddings:
    global _embeddings
    if _embeddings is None:
        _embeddings = OllamaEmbeddings(
            model="nomic-embed-text", base_url=settings.base_url
        )
    return _embeddings


def _get_store() -> InMemoryVectorStore:
    global _store
    if _store is None:
        if INDEX_PATH.exists():
            _store = InMemoryVectorStore.load(str(INDEX_PATH), _get_embeddings())
        else:
            _store = InMemoryVectorStore(_get_embeddings())
    return _store


def _read_file(path: Path) -> str:
    text = ""
    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return ""
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
    elif path.suffix.lower() == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            return ""
        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides, start=1):
            parts = [f"[Slide {i}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            parts.append(line)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [c.text for c in row.cells]
                        parts.append(" | ".join(cells))
            if len(parts) > 1:
                slides.append("\n".join(parts))
        text = "\n\n".join(slides)
    else:
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
    return text


def ingest_directory(directory: Path | None = None) -> dict:
    directory = directory or DOCS_DIR
    directory.mkdir(parents=True, exist_ok=True)
    files = [p for p in directory.iterdir() if p.suffix.lower() in {".txt", ".md", ".pdf", ".pptx", ".py", ".json", ".csv"}]
    docs = []
    for f in files:
        content = _read_file(f)
        if not content.strip():
            continue
        chunks = _splitter.split_text(content)
        for i, chunk in enumerate(chunks):
            docs.append(Document(page_content=chunk, metadata={"source": f.name, "chunk": i}))
    if not docs:
        return {"indexed_files": 0, "chunks": 0}
    store = _get_store()
    store.add_documents(docs)
    INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)
    store.dump(str(INDEX_PATH))
    return {"indexed_files": len(files), "chunks": len(docs)}


def search(query: str, k: int = 4) -> list[Document]:
    if not INDEX_PATH.exists():
        return []
    store = _get_store()
    hits = store.similarity_search(query, k=24)
    ql = query.lower()
    for word, ext in (("pdf", ".pdf"), ("ppt", ".pptx"), ("powerpoint", ".pptx"), ("slide", ".pptx")):
        if word in ql:
            filtered = [h for h in hits if ext in h.metadata.get("source", "").lower()]
            hits = filtered or hits
            break
    by_source: dict[str, list[Document]] = {}
    for h in hits:
        src = h.metadata.get("source", "?")
        if len(by_source.get(src, [])) < 3:
            by_source.setdefault(src, []).append(h)
    out = []
    while len(out) < k and any(by_source.values()):
        for src in list(by_source):
            if by_source[src]:
                out.append(by_source[src].pop(0))
    return out


def status() -> dict:
    files = []
    if DOCS_DIR.exists():
        files = [p.name for p in DOCS_DIR.iterdir() if p.is_file()]
    return {
        "docs_dir": str(DOCS_DIR),
        "files": files,
        "indexed": INDEX_PATH.exists(),
    }


def _filename_fallback(question: str) -> list[Document]:
    ql = question.lower()
    docs: list[Document] = []
    for p in sorted(DOCS_DIR.iterdir()) if DOCS_DIR.exists() else []:
        if not p.is_file() or p.name.lower() not in ql:
            continue
        text = _read_file(p)
        if not text.strip():
            continue
        for chunk in _splitter.split_text(text)[:6]:
            docs.append(Document(page_content=chunk, metadata={"source": p.name, "chunk": 0}))
    return docs


async def answer_from_docs(question: str) -> str:
    ql = question.lower()
    wanted_exts: list[str] = []
    if "pdf" in ql:
        wanted_exts.append(".pdf")
    if any(w in ql for w in ("ppt", "powerpoint", "slide")):
        wanted_exts.append(".pptx")
    wanted_sources: list[str] = []
    if DOCS_DIR.exists():
        for p in DOCS_DIR.iterdir():
            if p.is_file() and p.name.lower() in ql:
                wanted_sources.append(p.name.lower())
    if wanted_sources:
        store = _get_store()
        raw = store.similarity_search(question, k=200) if INDEX_PATH.exists() else []
        hits = [
            h for h in raw
            if h.metadata.get("source", "").lower() in wanted_sources
        ][:6]
        if not hits:
            hits = _filename_fallback(question)
    elif wanted_exts:
        store = _get_store()
        raw = store.similarity_search(question, k=200) if INDEX_PATH.exists() else []
        hits = [
            h for h in raw
            if any(e in h.metadata.get("source", "").lower() for e in wanted_exts)
        ][:6] or search(question, k=6)
    else:
        hits = search(question, k=6)
    if not hits:
        return "No documents indexed. Add files to data/docs and call /api/rag/ingest."
    context = "\n\n---\n\n".join(
        f"[{h.metadata['source']}]\n{h.page_content}" for h in hits
    )
    from langchain_core.messages import HumanMessage, SystemMessage

    from genflow.flows.research_flow import _script_hint

    messages = [
        SystemMessage(
            "You are GenFlow-AI answering questions about the user's uploaded "
            "documents. Rules: "
            "1) Answer using the provided context. If truly not in context, say "
            "you don't know. "
            "2) If the question is general (e.g. 'what is in this', 'summary', "
            "'is ppt me kya h'), summarize the key points from the context. "
            "3) Mention the source file name(s) used. "
            "4) Do NOT use emojis unless the user explicitly asks. "
            + _script_hint(question)
        ),
        HumanMessage(f"Context:\n{context}\n\nQuestion: {question}"),
    ]
    reply = await ainvoke_with_retry(get_llm(), messages)
    from genflow.flows.research_flow import force_roman_if_typed, strip_emojis

    return strip_emojis(force_roman_if_typed(question, reply.content))
